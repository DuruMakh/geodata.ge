"""Separate openpyxl source walk; never imports normalizer readers/helpers."""
import argparse
from collections import Counter, defaultdict
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
import openpyxl

ROOT=Path(__file__).resolve().parent
METADATA_FIELDS=('family','geography_id','partner_code','partner_label_en','source_group_id','source_group_label_en','classification','classification_level','product_code','product_label_en','dimension','group_id','group_label_en','geography_label_en','attribution_basis','service_id','service_label_en','year','flow','item_id','source_unit','value_status','publication_status','role','source_block','source_id','source_sheet','source_cell','source_label','source_number_format')

def digest(values):
    return hashlib.sha256(json.dumps(sorted(values),ensure_ascii=False,separators=(',',':')).encode('utf-8')).hexdigest()

def independent_code(value,classification):
    if classification=='sitc4':
        number=Decimal(str(value))
        return f'{number:05.1f}'
    width={'hs4':4,'hs6':6,'sitc1':1,'bec1':1}.get(classification,3)
    text=str(int(value)) if isinstance(value,(int,float)) else str(value).strip()
    return text.zfill(width)

def walk_sources(package_root: Path):
    manifest=json.loads((package_root/'full-source-manifest.json').read_text(encoding='utf-8'))
    if len(manifest)!=45 or len({r['source_id'] for r in manifest})!=45:
        raise ValueError('source_inventory: independent source count')
    tables=json.loads((package_root/'source-layouts.json').read_text(encoding='utf-8'))
    by_source=defaultdict(list)
    for table in tables: by_source[table['source_id']].append(table)
    for source in manifest:
        path=package_root/source['local_file']; data=path.read_bytes()
        if hashlib.sha256(data).hexdigest()!=source['sha256'] or len(data)!=source['bytes']:
            raise ValueError(f'source_fingerprint: {source["source_id"]}')
        if path.suffix!='.xlsx': continue
        with ZipFile(path) as zipped:
            if zipped.testzip() is not None: raise ValueError('source_fingerprint: independent ZIP check')
        book=openpyxl.load_workbook(path,read_only=True,data_only=True)
        try:
            for table in by_source[source['source_id']]:
                if table['source_sheet'] not in book.sheetnames: raise ValueError('source_layout: independent exact sheet')
                sheet=book[table['source_sheet']]
                needed={cell['cell'] for row in table['rows'] for cell in row['label_cells']+row['code_cells']}
                needed.update(h['cell'] for h in table['header_cells'])
                needed.update(f'{col}{r["row_index"]}' for r in table['rows'] for col in table['year_columns'].values())
                cells={c.coordinate:c for row in sheet.iter_rows() for c in row if getattr(c,'coordinate',None) in needed}
                for expected in table['header_cells']+[c for row in table['rows'] for c in row['label_cells']+row['code_cells']]:
                    cell=cells.get(expected['cell'])
                    text='' if cell is None or cell.value is None else str(cell.value)
                    fmt='General' if cell is None else cell.number_format
                    if text!=expected['value'] or fmt!=expected['format']:
                        raise ValueError(f'source_layout: independent metadata {source["source_id"]}:{expected["cell"]}')
                for row in table['rows']:
                    dimensions=dict(row['dimensions'])
                    if row['code_cells']:
                        native=cells[row['code_cells'][0]['cell']].value
                        if dimensions.get('classification'):
                            code=independent_code(native,dimensions['classification'])
                            if code!=dimensions['product_code']: raise ValueError('source_layout: independent product identity')
                            dimensions['product_code']=code
                        elif 'partner_code' in dimensions:
                            code=independent_code(native,'country')
                            if code!=dimensions['partner_code']: raise ValueError('source_layout: independent partner identity')
                            dimensions['partner_code']=code
                    label=' | '.join('' if cells.get(c['cell']) is None or cells[c['cell']].value is None else str(cells[c['cell']].value) for c in row['label_cells'])
                    for year,column in table['year_columns'].items():
                        coordinate=f'{column}{row["row_index"]}'; cell=cells.get(coordinate)
                        value=None if cell is None else cell.value
                        status='blank' if value is None else 'not_applicable' if value=='-' else 'numeric' if isinstance(value,(int,float)) and not isinstance(value,bool) else 'invalid'
                        if status=='invalid': raise ValueError('cell_type: independent monetary type')
                        metadata=dict(dimensions,family=table['family'],year=year,flow=table['flow'],item_id=row['item_id'],source_block=table['source_block'],source_id=source['source_id'],source_sheet=table['source_sheet'],source_cell=coordinate,source_unit=table['source_unit'],source_label=label,source_number_format='General' if cell is None else cell.number_format,value_status=status,publication_status='unspecified',role=row['role'])
                        yield {field:metadata.get(field,'') for field in METADATA_FIELDS},value
        finally:
            book.close()

def capture_inventory(package_root: Path) -> dict:
    groups=defaultdict(list)
    for row,value in walk_sources(package_root):
        key=tuple(row[f] for f in ('source_id','source_sheet','source_block','family','flow','year'))
        groups[key].append(row)
    blocks=[]
    for key,rows in sorted(groups.items()):
        block=dict(zip(('source_id','source_sheet','source_block','family','flow','year'),key))
        block.update(source_key_count=len(rows),key_sha256=digest([(r['source_id'],r['source_sheet'],r['source_cell']) for r in rows]),metadata_sha256=digest([sorted(r.items()) for r in rows]),identity_count=len({r['item_id'] for r in rows}),status_counts=dict(Counter(r['value_status'] for r in rows)))
        blocks.append(block)
    return {'full_manifest_sha256':hashlib.sha256((package_root/'full-source-manifest.json').read_bytes()).hexdigest(),'layout_sha256':hashlib.sha256((package_root/'source-layouts.json').read_bytes()).hexdigest(),'source_count':45,'workbook_count':31,'table_count':len(json.loads((package_root/'source-layouts.json').read_text(encoding='utf-8'))),'observation_count':sum(g['source_key_count'] for g in blocks),'blocks':blocks}

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--capture-inventory',action='store_true'); args=parser.parse_args()
    if not args.capture_inventory: parser.error('Full prepared verification is implemented in Task 5.')
    path=ROOT/'expected-observation-inventory.json'
    if path.exists(): raise ValueError('source_inventory: capture already exists; review changes explicitly')
    result=capture_inventory(ROOT)
    path.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k!='blocks'}))
