"""Audit this frozen Geostat source capture; no network or application writes."""
import argparse
import csv
import hashlib
import io
import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
TOLERANCE = 0.001  # Thousand USD: one dollar, for stored-value arithmetic noise.
ANNUAL = {
    'Export-Country_1995-2026.xlsx': [('1995-2025-years', 1995, 2025)],
    'Import-Country-1995-2026.xlsx': [('1995-2025-years', 1995, 2025)],
    'Export-_Country_Group-1995-2026.xlsx': [('1995-2025-years', 1995, 2025)],
    'Import_Country_Group-1995-2026.xlsx': [('1995-2025-years', 1995, 2025)],
    'Export-Product-by-4-digit-2015-2026.xlsx': [('2020-2025-years', 2020, 2025), ('2015-2019-years', 2015, 2019)],
    'Import-Product-by-4-digit-2015-2026.xlsx': [('2020-2025-years', 2020, 2025), ('2015-2019-years', 2015, 2019)],
    'Export-Product-by-4-digit-2000-2014.xlsx': [('2000-2014-years', 2000, 2014)],
    'Import-Product-by-4-digit-2000-2014.xlsx': [('2000-2014-years', 2000, 2014)],
    'Export-Product-by-4-digit-1995-1999.xlsx': [('1995-1999-years', 1995, 1999)],
    'Import-products--1995-1999_eng.xlsx': [('Import_products-1995-1999-years', 1995, 1999)],
    'Export-SITC_2000-2026.xlsx': [('SITC-2000-2025-years', 2000, 2025)],
    'Import-SITC_2000-2026.xlsx': [('SITC-2000-2025-years', 2000, 2025)],
    'Export-BEC_2000-2026.xlsx': [('BEC-2000-2025-years', 2000, 2025)],
    'Import-BEC-2000-2026.xlsx': [('BEC-2000-2025-years', 2000, 2025)],
    'Domestic-Export-Country_2014-2026.xlsx': [('2014-2025-years', 2014, 2025)],
    'Domestic-Exports_by-4-digit-2014-2026.xlsx': [('2020-2025-years', 2020, 2025), ('2015-2019-years', 2015, 2019), ('2014 year', 2014, 2014)],
}
SUPPORTING = {'export-by-regions-and-SITC-4-digital.xlsx', 'import-by-regions-and-SITC-4-digital.xlsx'}
EXPECTED_FILES = set(ANNUAL) | SUPPORTING | {'FTrade_1995-2026.xlsx'}

def year(value):
    text = str(value).strip()
    return int(text) if re.fullmatch(r'(?:199[5-9]|20[0-2]\d)', text) and int(text) <= 2025 else None

def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)

def audit():
    checks, tables, books = [], [], {}
    def check(kind, file, sheet, period, expected, actual, tolerance=0):
        passed = abs(expected - actual) <= tolerance if number(expected) and number(actual) else expected == actual
        checks.append(dict(check=kind, file=file, sheet=sheet, period=period, expected=expected, actual=actual, tolerance=tolerance, passed=passed))

    manifest = json.loads((ROOT / 'source-manifest.json').read_text(encoding='utf-8'))
    check('manifest_source_count', '', '', '', 27, len(manifest))
    check('unique_source_ids', '', '', '', len(manifest), len({r['source_id'] for r in manifest}))
    xlsx = {Path(r['local_file']).name for r in manifest if r['local_file'].endswith('.xlsx')}
    check('expected_workbook_inventory', '', '', '', sorted(EXPECTED_FILES), sorted(xlsx))
    for record in manifest:
        path = ROOT / record['local_file']
        data = path.read_bytes()
        check('file_sha256', path.name, '', '', record['sha256'], hashlib.sha256(data).hexdigest())
        check('file_bytes', path.name, '', '', record['bytes'], len(data))
        if path.suffix.lower() == '.xlsx':
            with ZipFile(path) as zipped:
                check('xlsx_zip_integrity', path.name, '', '', None, zipped.testzip())
            books[path.name] = load_workbook(path, read_only=True, data_only=True)
    check('worksheet_inventory', '', '', '', 55, sum(len(book.sheetnames) for book in books.values()))

    overview = list(books['FTrade_1995-2026.xlsx'].worksheets[0].iter_rows(values_only=True))
    check('overview_unit', 'FTrade_1995-2026.xlsx', '1995-2026', '', '(Mill. USD)', overview[2][0])
    overview_years = {year(value): i for i, value in enumerate(overview[3]) if year(value) is not None}
    check('overview_full_year_inventory', '', '', '', list(range(1995, 2026)), list(overview_years))
    baseline = {
        'export': {y: overview[4][col] * 1000 for y, col in overview_years.items()},
        'import': {y: overview[19][col] * 1000 for y, col in overview_years.items()},
    }
    domestic = books['Domestic-Export-Country_2014-2026.xlsx']['2014-2025-years']
    domestic_rows = list(domestic.iter_rows(values_only=True))
    baseline['domestic'] = {year(v): domestic_rows[4][i] for i, v in enumerate(domestic_rows[3]) if year(v) is not None}

    for flow, total_row, first_month in [('export', 4, 6), ('import', 19, 21)]:
        for col, value in enumerate(overview[3]):
            label = str(value).strip()
            if year(value) is None and label != '2026*':
                continue
            monthly = [overview[i][col] for i in range(first_month, first_month + 12) if number(overview[i][col])]
            check('national_month_count', 'FTrade_1995-2026.xlsx', '1995-2026', flow + ':' + label, 8 if label == '2026*' else 12, len(monthly))
            check('national_months_sum', 'FTrade_1995-2026.xlsx', '1995-2026', flow + ':' + label, overview[total_row][col] * 1000, sum(monthly) * 1000, TOLERANCE)

    for filename, expected_sheets in ANNUAL.items():
        flow = 'domestic' if filename.startswith('Domestic') else 'import' if filename.lower().startswith('import') else 'export'
        for sheet_name, start, end in expected_sheets:
            sheet = books[filename][sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            cols = {year(v): i for i, v in enumerate(rows[3]) if year(v) is not None}
            check('annual_year_inventory', filename, sheet_name, '', list(range(start, end + 1)), list(cols))
            check('annual_unit', filename, sheet_name, '', '(Thsd. USD)', rows[2][0])
            coded = []
            other = []
            for index, row in enumerate(rows[6:], 7):
                code, label = row[:2]
                if label is None:
                    continue
                if number(code) or (isinstance(code, str) and re.fullmatch(r'\d+', code.strip())):
                    coded.append((str(int(float(code))), index, row))
                elif str(label).strip() == 'Other commodities':
                    other.append((index, row))
            duplicates = [code for code, count in Counter(row[0] for row in coded).items() if count > 1]
            check('unique_annual_codes', filename, sheet_name, '', [], duplicates)
            check('domestic_other_row_inventory', filename, sheet_name, '', 1 if flow == 'domestic' and '4-digit' in filename else 0, len(other))
            for y, col in cols.items():
                total = rows[4][col]
                check('annual_total_cross_source', filename, sheet_name, y, baseline[flow][y], total, TOLERANCE)
                if coded:
                    detail_sum = sum(row[col] for _, _, row in coded if number(row[col])) + sum(row[col] for _, row in other if number(row[col]))
                    check('annual_detail_sum', filename, sheet_name, y, total, detail_sum, TOLERANCE)
            tables.append(dict(file=filename, sheet=sheet_name, start_year=start, end_year=end, coded_rows=len(coded), other_rows=[index for index, _ in other], total_cells=[f'{sheet.cell(5, col+1).coordinate}' for col in cols.values()]))
    for y, total in baseline['domestic'].items():
        check('domestic_within_total_exports', '', '', y, True, total <= baseline['export'][y] + TOLERANCE)
    for book in books.values():
        book.close()
    summary = dict(source_files=len(manifest), workbooks=len(xlsx), worksheets=55, annual_tables=len(tables), checks=len(checks), passed=sum(c['passed'] for c in checks), failed=sum(not c['passed'] for c in checks), tolerance_thousand_usd=TOLERANCE, check_counts=dict(Counter(c['check'] for c in checks)), annual_tables_reviewed=tables, limitations=['Annual totals and detail sums were checked; individual source cells were not independently re-extracted.', 'Only national monthly sums were checked; individual monthly series and quarterly regional data were not reconciled.', 'Country-group membership and historical product-code equivalence were not reviewed.', 'Domestic product rows are a selected list plus an explicit Other commodities remainder.', '2026 is January-August preliminary data, excluded from complete annual history.'])
    stream = io.StringIO(newline='')
    writer = csv.DictWriter(stream, fieldnames=list(checks[0]))
    writer.writeheader()
    writer.writerows(checks)
    return summary, stream.getvalue()

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--write', action='store_true', help='Write source-audit evidence after inspecting a new capture.')
    args = parser.parse_args()
    summary, csv_text = audit()
    json_text = json.dumps(summary, indent=2, ensure_ascii=False) + '\n'
    if args.write:
        (ROOT / 'validation.json').write_text(json_text, encoding='utf-8')
        (ROOT / 'reconciliation.csv').write_text(csv_text, encoding='utf-8-sig', newline='')
    else:
        assert (ROOT / 'validation.json').read_text(encoding='utf-8') == json_text, 'Validation report differs from source audit'
        with (ROOT / 'reconciliation.csv').open(encoding='utf-8-sig', newline='') as handle:
            assert handle.read() == csv_text, 'Reconciliation evidence differs from source audit'
    print(json.dumps({key: value for key, value in summary.items() if key not in ['annual_tables_reviewed', 'limitations']}, indent=2))
    if summary['failed']:
        raise SystemExit(1)
