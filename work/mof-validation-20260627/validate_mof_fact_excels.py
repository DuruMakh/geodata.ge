from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber
from pypdf import PdfReader


REPO_ROOT = Path(__file__).resolve().parents[2]
MOF_ROOT = REPO_ROOT / "docs" / "Raw Data" / "Expenditure" / "mof.ge"
FINAL_DIR = MOF_ROOT / "final-fact-files-2004-2025"
EXCEL_DIR = MOF_ROOT / "excel-fact-files-2004-2025"
MANIFEST = MOF_ROOT / "excel-fact-files-2004-2025-manifest.csv"
OLD_DIR = MOF_ROOT / "old"
WORK_DIR = Path(__file__).resolve().parent
OFFICE_DIR = WORK_DIR / "office-rerun"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(REPO_ROOT))
    except ValueError:
        return str(path)


def load_manifest() -> list[dict[str, str]]:
    with MANIFEST.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    return rows


def clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    text = " ".join(str(value).replace("\n", " ").split())
    return text if text != "" else None


def rows_from_xlsx(path: Path, sheet_name: str | None = None) -> list[list[Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        return [[clean_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def workbook_stats(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = []
        formulas = 0
        errors = 0
        non_empty = 0
        numeric_values: list[float] = []
        text_values: list[str] = []
        source_paths: list[str] = []

        for sheet in workbook.worksheets:
            max_row = sheet.max_row or 0
            max_col = sheet.max_column or 0
            sheet_non_empty = 0
            for row in sheet.iter_rows(values_only=False):
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    sheet_non_empty += 1
                    non_empty += 1
                    if isinstance(value, str) and value.startswith("="):
                        formulas += 1
                    if isinstance(value, str) and value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                        errors += 1
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        numeric_values.append(float(value))
                    elif isinstance(value, str):
                        text = clean_cell(value)
                        if text:
                            text_values.append(text)
                            if "docs" in text and ("mof.ge" in text or "Raw Data" in text):
                                source_paths.append(text)
            sheets.append(
                {
                    "name": sheet.title,
                    "max_row": max_row,
                    "max_column": max_col,
                    "non_empty_cells": sheet_non_empty,
                }
            )

        rounded_numeric = [round(v, 6) for v in numeric_values if math.isfinite(v)]
        return {
            "sheets": sheets,
            "sheet_count": len(workbook.worksheets),
            "non_empty_cells": non_empty,
            "formula_cells": formulas,
            "error_literals": errors,
            "numeric_cell_count": len(numeric_values),
            "numeric_sum": round(sum(rounded_numeric), 6),
            "numeric_min": min(rounded_numeric) if rounded_numeric else None,
            "numeric_max": max(rounded_numeric) if rounded_numeric else None,
            "text_cell_count": len(text_values),
            "source_path_cells": source_paths[:10],
        }
    finally:
        workbook.close()


def pdf_table_rows(path: Path) -> tuple[list[list[Any]], int]:
    rows: list[list[Any]] = []
    max_columns = 0
    with pdfplumber.open(path) as pdf:
        for page_index, page in enumerate(pdf.pages, start=1):
            for table_index, table in enumerate(page.extract_tables(), start=1):
                for row_index, row in enumerate(table, start=1):
                    values = [clean_cell(cell) for cell in row]
                    max_columns = max(max_columns, len(values))
                    rows.append([page_index, table_index, row_index, *values])
    return rows, max_columns


def compare_matrix(left: list[list[Any]], right: list[list[Any]]) -> dict[str, Any]:
    max_rows = max(len(left), len(right))
    mismatch_examples = []
    mismatch_count = 0

    for index in range(max_rows):
        left_row = left[index] if index < len(left) else None
        right_row = right[index] if index < len(right) else None
        if left_row != right_row:
            mismatch_count += 1
            if len(mismatch_examples) < 10:
                mismatch_examples.append(
                    {
                        "row_index_1_based": index + 1,
                        "expected": left_row,
                        "actual": right_row,
                    }
                )

    return {
        "expected_rows": len(left),
        "actual_rows": len(right),
        "mismatch_count": mismatch_count,
        "mismatch_examples": mismatch_examples,
    }


def normalized_csv_rows(path: Path) -> list[list[Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [[clean_cell(value) for value in row] for row in csv.reader(handle)]


def text_number_tokens(text: str) -> Counter[str]:
    tokens = re.findall(r"(?<!\w)-?\d[\d\s,.\u00a0]*\d|(?<!\w)\d(?!\w)", text)
    cleaned = []
    for token in tokens:
        value = token.replace("\u00a0", " ").strip()
        value = re.sub(r"\s+", "", value)
        value = value.replace(",", ".")
        if value:
            cleaned.append(value)
    return Counter(cleaned)


def workbook_number_tokens(path: Path) -> Counter[str]:
    counter: Counter[str] = Counter()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        as_float = float(value)
                        if as_float.is_integer():
                            counter[str(int(as_float))] += 1
                        else:
                            counter[(f"{as_float:.6f}").rstrip("0").rstrip(".")] += 1
                    elif isinstance(value, str):
                        counter.update(text_number_tokens(value))
    finally:
        workbook.close()
    return counter


def pdf_text_probe(path: Path, output_xlsx: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    pages = []
    combined_text = ""
    for index, page in enumerate(reader.pages[:3], start=1):
        text = page.extract_text() or ""
        pages.append({"page": index, "chars": len(text), "snippet": " ".join(text.split())[:400]})
        combined_text += "\n" + text
    pdf_numbers = text_number_tokens(combined_text)
    xlsx_numbers = workbook_number_tokens(output_xlsx)
    common = sum((pdf_numbers & xlsx_numbers).values())
    return {
        "first_pages": pages,
        "pdf_number_tokens_first_3_pages": sum(pdf_numbers.values()),
        "xlsx_number_tokens": sum(xlsx_numbers.values()),
        "common_number_tokens": common,
        "year_markers_in_first_pages": sorted(set(re.findall(r"20\d{2}", combined_text))),
    }


def old_hash_index() -> dict[str, list[str]]:
    index: dict[str, list[str]] = defaultdict(list)
    if not OLD_DIR.exists():
        return index
    for path in OLD_DIR.rglob("*"):
        if not path.is_file():
            continue
        index[sha256(path)].append(rel(path))
    return index


def inspect_metadata_source(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "metadata" not in workbook.sheetnames:
            return {"metadata_sheet": False, "rows": []}
        rows = []
        for row in workbook["metadata"].iter_rows(values_only=True):
            rows.append([clean_cell(value) for value in row])
        return {"metadata_sheet": True, "rows": rows}
    finally:
        workbook.close()


def main() -> int:
    rows = load_manifest()
    old_index = old_hash_index()

    report: dict[str, Any] = {
        "root": str(REPO_ROOT),
        "manifest": rel(MANIFEST),
        "techniques": [
            "coverage_and_manifest_consistency",
            "sha256_exact_copy_for_copied_xlsx",
            "old_folder_hash_match_for_source_provenance",
            "xlsx_openability_shape_numeric_integrity",
            "pdf_table_reextract_vs_xlsx_matrix",
            "word_table_rerun_csv_vs_xlsx_matrix_when_available",
            "xls_to_xlsx_rerun_vs_output_matrix_when_available",
            "pdf_text_numeric_token_probe",
        ],
        "coverage": {},
        "years": {},
        "summary": {},
    }

    years = [int(row["year"]) for row in rows]
    report["coverage"] = {
        "manifest_row_count": len(rows),
        "years": years,
        "expected_years": list(range(2004, 2026)),
        "missing_years": sorted(set(range(2004, 2026)) - set(years)),
        "duplicate_years": sorted(year for year, count in Counter(years).items() if count > 1),
    }

    for row in rows:
        year = int(row["year"])
        source = Path(row["source_file"])
        output = Path(row["output_file"])
        method = row["method"]
        source_hash = sha256(source) if source.exists() else None
        output_hash = sha256(output) if output.exists() else None

        year_report: dict[str, Any] = {
            "source": rel(source),
            "output": rel(output),
            "method": method,
            "source_exists": source.exists(),
            "output_exists": output.exists(),
            "source_sha256": source_hash,
            "output_sha256": output_hash,
            "source_hash_matches_in_old": old_index.get(source_hash, [])[:10] if source_hash else [],
            "checks": {},
        }

        if output.exists():
            year_report["workbook_stats"] = workbook_stats(output)

        if method == "copied_xlsx":
            year_report["checks"]["copy_hash_match"] = source_hash == output_hash
        else:
            year_report["checks"]["copy_hash_match"] = None

        if method == "pdfplumber_tables_to_xlsx" and source.exists() and output.exists():
            pdf_rows, max_columns = pdf_table_rows(source)
            headers = ["source_page", "source_table", "row_index"] + [
                f"col_{index}" for index in range(1, max_columns + 1)
            ]
            expected = [headers] + [row + [None] * (len(headers) - len(row)) for row in pdf_rows]
            actual = rows_from_xlsx(output, "extracted_tables")
            year_report["checks"]["pdf_table_matrix_match"] = compare_matrix(expected, actual)
            year_report["checks"]["pdf_metadata"] = inspect_metadata_source(output)
            year_report["checks"]["pdf_text_probe"] = pdf_text_probe(source, output)

        if method == "word_com_tables_csv_to_xlsx" and output.exists():
            csv_path = OFFICE_DIR / f"{year}-word-tables-rerun.csv"
            if csv_path.exists():
                expected = normalized_csv_rows(csv_path)
                actual = rows_from_xlsx(output, "extracted_tables")
                year_report["checks"]["word_table_rerun_matrix_match"] = compare_matrix(expected, actual)
            else:
                year_report["checks"]["word_table_rerun_matrix_match"] = {
                    "skipped": f"Missing {csv_path}",
                }
            year_report["checks"]["word_metadata"] = inspect_metadata_source(output)

        if method == "excel_com_xls_to_xlsx" and output.exists():
            rerun_path = OFFICE_DIR / f"{year}-xls-rerun.xlsx"
            if rerun_path.exists():
                expected = rows_from_xlsx(rerun_path)
                actual = rows_from_xlsx(output)
                year_report["checks"]["xls_rerun_matrix_match"] = compare_matrix(expected, actual)
                year_report["checks"]["xls_rerun_sha256"] = sha256(rerun_path)
            else:
                year_report["checks"]["xls_rerun_matrix_match"] = {
                    "skipped": f"Missing {rerun_path}",
                }

        report["years"][str(year)] = year_report

    failures: list[str] = []
    unresolved: list[str] = []
    for year_text, year_report in report["years"].items():
        stats = year_report.get("workbook_stats", {})
        if not year_report["source_exists"]:
            failures.append(f"{year_text}: source missing")
        if not year_report["output_exists"]:
            failures.append(f"{year_text}: output missing")
        if stats and stats["non_empty_cells"] == 0:
            failures.append(f"{year_text}: output workbook is empty")
        if stats and stats["error_literals"] > 0:
            failures.append(f"{year_text}: output workbook contains error literals")
        if year_report["method"] == "copied_xlsx" and not year_report["checks"].get("copy_hash_match"):
            failures.append(f"{year_text}: copied_xlsx hash mismatch")
        pdf_check = year_report["checks"].get("pdf_table_matrix_match")
        if isinstance(pdf_check, dict) and pdf_check.get("mismatch_count", 0) != 0:
            failures.append(f"{year_text}: PDF table extraction mismatch")
        word_check = year_report["checks"].get("word_table_rerun_matrix_match")
        if isinstance(word_check, dict) and word_check.get("mismatch_count", 0) != 0:
            failures.append(f"{year_text}: Word rerun extraction mismatch")
        xls_check = year_report["checks"].get("xls_rerun_matrix_match")
        if isinstance(xls_check, dict) and xls_check.get("mismatch_count", 0) != 0:
            failures.append(f"{year_text}: XLS rerun conversion mismatch")
        if not year_report["source_hash_matches_in_old"]:
            unresolved.append(f"{year_text}: no exact old-folder hash match for source file")

    report["summary"] = {
        "failure_count": len(failures),
        "failures": failures,
        "unresolved_provenance_notes": unresolved,
        "years_checked": len(report["years"]),
    }

    output_json = WORK_DIR / "mof_fact_excel_validation_report.json"
    output_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(report["coverage"], ensure_ascii=False, indent=2))
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))
    print(f"report={output_json}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
